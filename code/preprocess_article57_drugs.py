#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import load_workbook


OUTPUT_FILES = {
    "drug_type": "drug_type.csv",
    "drug_support_phone": "drug_support_phone.csv",
    "substances": "substances.csv",
    "has_substances": "has_substances.csv",
}
LEGACY_FILES = ("DB_DrugType.csv", "DB_DrugSupportPhone.csv")
PHONE_PATTERN = re.compile(r"(?:\+|00)?\d[\d\s().-]{7,}\d")
SALT_SUFFIXES = {
    "acetate", "anhydrous", "base", "besilate", "bitartrate", "bromide", "calcium", "chloride",
    "citrate", "dihydrate", "diphosphate", "fumarate", "hydrate", "hydrobromide", "hydrochloride",
    "hydrated", "iodide", "lactate", "magnesium", "maleate", "mesilate", "monohydrate", "nitrate", "oxalate",
    "phosphate", "potassium", "sodium", "succinate", "sulfate", "sulphate", "tartrate", "tosylate",
}
CATION_BASES = {
    "aluminium", "aluminum", "ammonium", "calcium", "chromium", "copper", "ferric", "ferrous",
    "lithium", "magnesium", "manganese", "potassium", "sodium", "zinc",
}
NON_STRIPPABLE_BASES = {
    "diethyl", "dimethyl", "disodium", "ethyl", "magnesium", "monosodium", "potassium", "sodium",
}
METADATA_AFTER_COMMA = re.compile(
    r"^(?:"
    r"adjusted|calculated|corresponding|consisting|containing|derived|equivalent|expressed|including|measured|"
    r"nominal|quantified|refined|standardi[sz]ed|extract(?:ing)?|extraction|solvent|ethanol|methanol|acetone|"
    r"water|planta|flos|folium|fructus|radix|rhizom|semen|cortex|herba|rec\.|aqu\.|cum|der\b|"
    r"ekstrahent|gekstrahent|etanol|metanol|"
    r"activated|anhydrous|dried|ep|hydrated|hydrous|medicinal|strain|synthetic|virgin|"
    r"disodium|monosodium"
    r")\b",
    re.IGNORECASE,
)


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.lower() == "nan":
        return ""
    return re.sub(r"\s+", " ", text.replace("\ufeff", " ")).strip()


def default_source_path(project_root: Path) -> Path:
    source = project_root / "data" / "raw_data" / "article-57-product-data_en.xlsx"
    if not source.exists():
        raise FileNotFoundError(f"Article 57 workbook not found: {source}")
    return source


def normalized_header(value: object) -> str:
    return clean(value).lower()


def find_column(headers: list[str], *needles: str) -> int:
    for idx, header in enumerate(headers):
        if all(needle in header for needle in needles):
            return idx
    raise ValueError(f"Could not find a workbook column containing: {', '.join(needles)}")


def find_header_row(worksheet) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        headers = [normalized_header(value) for value in row]
        if not any("product name" in header for header in headers):
            continue
        if not any("active substance" in header for header in headers):
            continue

        return row_number, {
            "product": find_column(headers, "product name"),
            "substance": find_column(headers, "active substance"),
            "route": find_column(headers, "route of administration"),
            "auth_country": find_column(headers, "product authorisation country"),
            "auth_holder": find_column(headers, "marketing authorisation holder"),
            "master_file": find_column(headers, "master file location"),
            "email": find_column(headers, "email address"),
            "phone": find_column(headers, "telephone number"),
        }
    raise ValueError(f"Could not locate the Article 57 header row in sheet {worksheet.title!r}")


def cell(row: tuple[object, ...], idx: int) -> object:
    return row[idx] if idx < len(row) else ""


def nearest_nonspace(text: str, start: int, step: int) -> str:
    idx = start
    while 0 <= idx < len(text):
        if not text[idx].isspace():
            return text[idx]
        idx += step
    return ""


def is_top_level_comma_separator(text: str, idx: int) -> bool:
    previous = nearest_nonspace(text, idx - 1, -1)
    following = nearest_nonspace(text, idx + 1, 1)
    if previous.isdigit() and following.isdigit():
        return False
    if idx > 0 and idx + 1 < len(text) and not text[idx - 1].isspace() and not text[idx + 1].isspace() and previous.isalpha() and following.isalpha():
        return False

    tail = clean(text[idx + 1:]).lower()
    if METADATA_AFTER_COMMA.match(tail):
        return False
    if re.match(r"^(?:and|or)\b", tail, flags=re.IGNORECASE):
        return False
    previous_segment = re.split(r"[,|]", text[:idx])[-1]
    if re.search(r"\bstrain\b", previous_segment, flags=re.IGNORECASE) and re.match(r"^[a-z0-9]{1,4}\b", tail, flags=re.IGNORECASE):
        return False
    if re.match(r"^[bc]\b\s*(?:and\b|$)", tail, flags=re.IGNORECASE):
        return False
    return True


def split_top_level(value: object) -> list[str]:
    text = clean(value)
    parts: list[str] = []
    buffer: list[str] = []
    depth = 0
    for idx, char in enumerate(text):
        if char in "([{":
            depth += 1
        elif char in ")]}" and depth > 0:
            depth -= 1

        if depth == 0 and (char == "|" or (char == "," and is_top_level_comma_separator(text, idx))):
            part = clean("".join(buffer))
            if part:
                parts.append(part)
            buffer = []
        else:
            buffer.append(char)

    part = clean("".join(buffer))
    if part:
        parts.append(part)
    return parts


def strip_trailing_salt(name: str) -> str:
    while True:
        words = name.split()
        if len(words) < 2:
            return name

        suffix = words[-1].rstrip(".").lower()
        base = " ".join(words[:-1]).strip()
        if suffix not in SALT_SUFFIXES or base.lower() in CATION_BASES or base.lower() in NON_STRIPPABLE_BASES:
            return name
        name = base


def normalize_substance_case(name: str) -> str:
    def fix_word(match: re.Match[str]) -> str:
        word = match.group(0)
        if len(word) > 1 and word[0].islower() and word[1].islower():
            return word[0].upper() + word[1:]
        return word

    return re.sub(r"\b[^\W\d_][^\W\d_']*\b", fix_word, name)


def minimize_substance_name(value: str) -> str:
    name = clean(value)
    if re.match(r"^(?:Extraction\s+Solvents?|Extraction\s+Agent|Extracton\s+Solvent|Extration\s+Solvent)\b", name, flags=re.IGNORECASE):
        return ""
    if re.search(r"\bWater\s+For\s+Injections?\b", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"\([^)]*\)", name):
        return ""
    if re.fullmatch(r"[\[\(].*[\]\)]", name) and re.search(r"\b(?:hab|urt\.|ethanol|methanol|v\.|m\s*/\s*m)\b", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^(?:Der|Drug Extract Ratio)\b", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^First\s+(?:Extraction\s+Solvents?|Water)\b", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^(?:And|Or)\b", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^Equivalent\s+To\b", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^(?:Auszugsmittel|Ekstrahent|Gekstrahent|Extract(?:ing)? Agent|Extract(?:ing)? Solvent|Extraction Agent|Extraction Liquid|Extraction Medium|Distillation Agent|Distillation Liquid)\b", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^(?:Composed Of|Consists Of|100 Ml Consists Of)\b", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"(?:E\.?P\.?|Hydrated|Powdered|Purified|Sterile|Water|Purified\s+Water|Sterile\s+Water|Water\s*,\s*Purified)", name, flags=re.IGNORECASE):
        return ""
    if re.match(r"^\[\d+[A-Za-z]+\]", name):
        name = re.sub(r"^\[\d+[A-Za-z]+\]\s*", "", name)
        name = name[:1].upper() + name[1:]
    name = re.sub(r"\s*\[[^\]]*(?:hab|ethanol|methanol|v\s*/\s*v|m\s*/\s*m)[^\]]*\]\s*", " ", name, flags=re.IGNORECASE)
    name = re.sub(r"^\((?:[0-9rsRS,]+|[±]|[dDlL]+)\)-", "", name)
    name = re.sub(
        r"^\(([^)]*[A-Za-z][^)]*)\)-",
        lambda match: re.sub(r"\s*,\s*", " ", match.group(1)) + " ",
        name,
    )
    name = re.sub(r"^(?:L|D|DL)-(?=[A-Za-z])", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*,\s*(?:activated|anhydrous|dried|e\.?p\.?|hydrated|hydrous|medicinal|powdered|purified|synthetic|virgin|disodium|monosodium)\s*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*\((?:D|L|DL)-[^)]*\)\s*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*,\s*(?:adjusted|calculated|corresponding|consisting|containing|equivalent|expressed|including|measured|nominal|quantified|refined|standardi[sz]ed)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*,\s*(?:extract(?:ing)?|extraction|solvent|ethanol|methanol|acetone|water|distillation)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*,\s*(?:planta|flos|folium|fructus|radix|rhizom|semen|cortex|herba|rec\.|aqu\.|cum)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\b(?:Ethanol|Etanol)\s+(Fluid|Dry)\s+Extract\b", r"\1 Extract", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*:\s*(?:Extract(?:ing)? Agent|Extract(?:ing)? Solvent|Extraction Agent|Extraction Solvents?|Extraction Medium|Extraction Liquid|Distillation Agent|Distillation Liquid|Auszugsmittel|Ekstrahent|Gekstrahent|Extracton Solvent|Extration Solvent)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*:\s*(?:Ethanol|Etanol|Methanol|Metanol)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*:\s*Water\s*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+(?:Extracting Agent|Extracting Solvent|Extraction Agent|Extraction Solvents?|Extraction Medium|Extraction Liquid|Distillation Agent|Distillation Liquid|Auszugsmittel|Ekstrahent|Gekstrahent|Extracton Solvent|Extration Solvent)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Ethanol\.\s*(?:Decoctum|Digestio|Extr\.?|Infusum).*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Rec\.\s+(?:Ethanol|Etanol)\.\s*(?:Digest\.?|Digestio|Extr\.?|Infusum).*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Ethanol\s+\d+(?:[.,]\d+)?%\s+And\s+Water.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+(?:And\s+|With\s+)(?:Ethanol|Etanol)\s+\d+(?:[.,]\d+)?\s*%.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+(?:Fe|Te)\s+With\s+Ethanol(?:\s*/\s*Ethanol)?[- ]Water.*$", "", name, flags=re.IGNORECASE)
    if re.match(r"^Rts,s\b", name, flags=re.IGNORECASE):
        name = "Rts,S Antigen"
    name = re.sub(
        r"^Preparation Of Fresh Herb Of Symphytum X Uplandicum\b.*$",
        "Symphytum X Uplandicum Fresh Herb Preparation",
        name,
        flags=re.IGNORECASE,
    )
    name = re.sub(r"\b(?:Ph\.?\s*Eur\.?|B\.?P\.?|U\.?S\.?P\.?)\b\.?", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*=\s*D\s*\d+\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bSolution\s*\(Prepared\s+From\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bAquos\.?(?:\s*\([^)]*\))?.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*-\s*Adsorbed\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bAdsorbed\s+On\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bOn\s+Aluminium\s+Hydroxide\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bProduced\s+In\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bBy\s+Rdna\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+With\s+A\s+Specific\s+Composition\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Measured\s+As\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bExtraction\s+(?:Agent|Solvents?)\s*:\s*", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*(?:,|:|\()\s*Der(?:\s+(?:Native|Genuine))?\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Der\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*\((?:Extraction\s+Solvents?|Extracting\s+Solvent|Extraction\s+Agent|Extraction\s+Medium|Extraction\s+Liquid|Distillation\s+Agent|Distillation\s+Liquid)\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+And\s+(?:Purified|Sterile)\s+Water\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Equivalent\s+To\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Corresponding\s+To\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Adjusted\s+To\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Calculated\s+As\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Expressed\s+As\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Containing\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(
        r"\s*\([^)]*(?:V/V|W/W|W/V|M/M|water|up to|parts|strain|extraction solvent|solvent|auszugsmittel|ratio|form|Hab|live,\s*attenuated|Der\b|Dev\b)[^)]*\)\s*",
        " ",
        name,
        flags=re.IGNORECASE,
    )
    name = re.sub(r"\s*\((?:\d+[.,:]?\d*\s*(?:-|–|:)?\s*)+\)\s*", " ", name)
    name = re.sub(r"\s*\([^)]*(?:/|:)\s*[^)]*\)\s*\d+\s*:\s*\d+\s*$", "", name)
    name = re.sub(r"\bDil\.\s*D\s*[0-9]+\b", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bSpag\.\s*[^,]*?\s*Dil\.\s*D\s*[0-9]+\b", "", name, flags=re.IGNORECASE)
    name = re.sub(r"^(?:Anhydrous|Hydrated)\s+", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Synthetic$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"^\d+(?:[.,]\d+)?(?:\s*(?:-|–|to)\s*\d+(?:[.,]\d+)?)?\s*(?:mg|g|ml|mcg|µg|iu|%)\s+of\s+", "", name, flags=re.IGNORECASE)
    name = re.sub(r"^\d+(?:[.,]\d+)?(?:\s*(?:-|–|to)\s*\d+(?:[.,]\d+)?)?\s*(?:mg|g|ml|mcg|µg|iu|%)\s+", "", name, flags=re.IGNORECASE)
    name = re.sub(r"^\d+(?:[.,]\d+)?\s*(?:p\.?|cz\.?)\s*", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+\d+(?:[.,]\d+)?\s*(?:p\.?|cz\.?)\s*(?:-|–).*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+\d+(?:[.,]\d+)?(?:\s*(?:-|–)\s*\d+(?:[.,]\d+)?)?\s*(?:p\.?|cz\.?|mg|g|ml|mcg|µg|iu|%)\.?$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+E\.?P\.?$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\b(Dry|Liquid|Soft)\s+Extract\s+Water\b", r"\1 Extract", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*(?:,|:|\()\s*Der(?:\s+(?:Native|Genuine))?\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+Der\b.*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"^.*\bEx:\s*", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s*/\s*", " / ", name)
    name = strip_trailing_salt(clean(name).strip(" ,;-"))
    if re.match(r"^(?:And|Or)\b", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"(?:E\.?P\.?|Hydrated|Powdered|Purified|Sterile|Synthetic|Water|Purified\s+Water|Sterile\s+Water|Water\s*,\s*Purified)", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"\([^)]*\)", name):
        return ""
    if re.fullmatch(r"[\d\s.,:;+\-/–%]+(?:p\.|cz\.)?", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"(?:[A-Za-z]|[0-9]+[A-Za-z]{1,3}|D\s*[0-9]*)", name):
        return ""
    if re.match(r"^(?:composed|corresponding|equivalent|calculated|expressed|containing|consisting|including|adjusted|measured|derived from)\b", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"(?:ethanol|methanol|acetone|water for injections?)(?:\s+[0-9.,]+\s*%)?", name, flags=re.IGNORECASE):
        return ""
    if re.fullmatch(r"(?:ethanol|methanol|acetone)(?:\s*\([^)]*\))?", name, flags=re.IGNORECASE):
        return ""
    if re.search(r"\b(?:O\.?\s*P\.?|Cz\.)\b", name, flags=re.IGNORECASE):
        return ""
    return normalize_substance_case(clean(name).strip(" ,;-."))


def split_substances(value: object) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for part in split_top_level(value):
        substance = minimize_substance_name(part)
        if substance and substance not in seen:
            seen.add(substance)
            out.append(substance)
    return out


def extract_phones(value: object) -> list[str]:
    text = re.sub(r"\b24\s*/\s*7\b", " ", clean(value), flags=re.IGNORECASE)
    seen: set[str] = set()
    phones: list[str] = []
    for match in PHONE_PATTERN.finditer(text):
        digits = re.sub(r"\D", "", match.group(0))
        if 10 <= len(digits) <= 14 and digits not in seen:
            seen.add(digits)
            phones.append(digits)
    return phones


def write_csv_atomic(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    tmp_path = path.with_name(f"{path.name}.tmp")
    with tmp_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    tmp_path.replace(path)


def extract_article57(source_path: Path, sheet_name: str | None = None) -> dict[str, list[dict]]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    header_row, columns = find_header_row(worksheet)

    drug_rows: list[dict] = []
    phone_rows: list[dict] = []
    substance_rows: list[dict] = []
    has_substance_rows: list[dict] = []
    substance_id_by_name: dict[str, int] = {}

    drug_id = 1
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number <= header_row:
            continue

        values = [clean(cell(row, columns[column])) for column in columns]
        if not any(values):
            continue

        drug_rows.append({
            "DrugID": drug_id,
            "Name": clean(cell(row, columns["product"])),
            "Route": clean(cell(row, columns["route"])),
            "AuthCountry": clean(cell(row, columns["auth_country"])),
            "AuthHolder": clean(cell(row, columns["auth_holder"])),
            "MasterFileLocation": clean(cell(row, columns["master_file"])),
            "Email": clean(cell(row, columns["email"])),
        })

        for phone in extract_phones(cell(row, columns["phone"])):
            phone_rows.append({"DrugID": drug_id, "Phone": phone})

        seen_substances_for_drug: set[int] = set()
        for substance in split_substances(cell(row, columns["substance"])):
            sub_id = substance_id_by_name.get(substance)
            if sub_id is None:
                sub_id = len(substance_id_by_name) + 1
                substance_id_by_name[substance] = sub_id
                substance_rows.append({
                    "ID": sub_id,
                    "Name": substance,
                })
            if sub_id in seen_substances_for_drug:
                continue
            seen_substances_for_drug.add(sub_id)
            has_substance_rows.append({"SubID": sub_id, "DrugID": drug_id})

        drug_id += 1

    if not drug_rows:
        raise RuntimeError(f"No Article 57 product rows were extracted from {source_path}")
    if not substance_rows:
        raise RuntimeError(f"No Article 57 substances were extracted from {source_path}")

    return {
        "drug_type": drug_rows,
        "drug_support_phone": phone_rows,
        "substances": substance_rows,
        "has_substances": has_substance_rows,
    }


def remove_legacy_files(output_dir: Path) -> None:
    for filename in LEGACY_FILES:
        path = output_dir / filename
        if path.exists():
            path.unlink()


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description="Extract Article 57 drug reference CSVs for HospitalDB.")
    parser.add_argument("--project-root", default=str(project_root))
    parser.add_argument("--source", default=None, help="Path to article-57-product-data_en.xlsx.")
    parser.add_argument("--output-dir", default=None, help="Output directory. Defaults to <project-root>/data.")
    parser.add_argument("--sheet", default=None, help="Workbook sheet name. Defaults to the active sheet.")
    parser.add_argument("--keep-legacy", action="store_true", help="Do not delete DB_DrugType.csv/DB_DrugSupportPhone.csv.")
    args = parser.parse_args()

    root = Path(args.project_root).expanduser().resolve()
    source = Path(args.source).expanduser().resolve() if args.source else default_source_path(root)
    output_dir = Path(args.output_dir).expanduser().resolve() if args.output_dir else root / "data"
    output_dir.mkdir(parents=True, exist_ok=True)

    extracted = extract_article57(source, args.sheet)
    write_csv_atomic(output_dir / OUTPUT_FILES["drug_type"], [
        "DrugID", "Name", "Route", "AuthCountry", "AuthHolder", "MasterFileLocation", "Email",
    ], extracted["drug_type"])
    write_csv_atomic(output_dir / OUTPUT_FILES["drug_support_phone"], ["DrugID", "Phone"], extracted["drug_support_phone"])
    write_csv_atomic(output_dir / OUTPUT_FILES["substances"], ["ID", "Name"], extracted["substances"])
    write_csv_atomic(output_dir / OUTPUT_FILES["has_substances"], ["SubID", "DrugID"], extracted["has_substances"])

    if not args.keep_legacy:
        remove_legacy_files(output_dir)

    print(f"Wrote {len(extracted['drug_type'])} drug types to {output_dir / OUTPUT_FILES['drug_type']}")
    print(f"Wrote {len(extracted['drug_support_phone'])} drug support phones to {output_dir / OUTPUT_FILES['drug_support_phone']}")
    print(f"Wrote {len(extracted['substances'])} substances to {output_dir / OUTPUT_FILES['substances']}")
    print(f"Wrote {len(extracted['has_substances'])} drug-substance links to {output_dir / OUTPUT_FILES['has_substances']}")


if __name__ == "__main__":
    main()
